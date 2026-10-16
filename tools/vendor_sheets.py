"""vendor_sheets.py: one order sheet per seller, in the team's purchase-form layout (run from repo root)."""
import csv, os, re
from copy import copy
from openpyxl import Workbook, load_workbook

SRC = "docs/full-order-sheet.csv"
OUT = "docs/vendor-orders"
# Blank copy of the team's purchase form; every cell style is taken from it
TPL = load_workbook(os.path.join(os.path.dirname(__file__), "vendor-order-template.xlsx")).active
BOARD = {"Drone parts (1 drone)": "Drone", "Flight controller parts (1 board)": "Flight controller", "Charging board parts (1 board)": "Charging board"}

rows = [r for r in csv.reader(open(SRC, encoding="utf-8-sig")) if r and r[1] and r[0] != "Group"]
sellers = []
for r in rows:
    if r[1] not in sellers:
        sellers.append(r[1])

def sku(r):
    m = re.search(r"amazon\.com/dp/([A-Z0-9]{10})", r[5])
    if m:
        return m.group(1)
    return "" if "/s?k=" in r[5] else r[3]

def item(r):
    # Drone parts are named by product; electronic parts read better by what they are
    return r[3] if r[0].startswith("Drone") else r[4]

def specs(r):
    keep = [s.strip() for s in re.split(r"(?<=\.)\s+", r[8]) if s.strip()]
    keep = [s for s in keep if not s.startswith(("Buy ", "Was ", "Take the price", "Provisional"))]
    text = " ".join(keep[:2])
    extra = r[4] if r[0].startswith("Drone") else ""
    return "; ".join(x for x in [BOARD.get(r[0], ""), extra, text] if x)

def style(dst, src):
    dst.font = copy(src.font); dst.fill = copy(src.fill); dst.border = copy(src.border)
    dst.alignment = copy(src.alignment); dst.number_format = src.number_format; dst.protection = copy(src.protection)

def fill_sheet(ws, seller_rows):
    for col in "ABCDEFGH":
        c = ws[col + "1"]; c.value = TPL[col + "1"].value; style(c, TPL[col + "1"])
    for i, r in enumerate(seller_rows, start=2):
        vals = [item(r), int(r[2]), float(r[6]) if r[6] else None, "=B%d*C%d" % (i, i), sku(r), specs(r), r[5], "=D%d" % i]
        for col, v in zip("ABCDEFGH", vals):
            c = ws["%s%d" % (col, i)]; c.value = v; style(c, TPL[col + "2"])
        ws["G%d" % i].hyperlink = r[5]
    n = len(seller_rows) + 2
    for k, (rr, val) in enumerate([(30, "Please select fastest shipping available"), (31, None)]):
        for col in "GH":
            c = ws["%s%d" % (col, n + k)]; c.value = TPL["%s%d" % (col, rr)].value; style(c, TPL["%s%d" % (col, rr)])
    ws["H%d" % (n + 1)] = "=SUM(H2:H%d)" % n
    for k, d in TPL.column_dimensions.items():
        ws.column_dimensions[k].width = d.width
    ws.row_dimensions[1].height = TPL.row_dimensions[1].height

os.makedirs(OUT, exist_ok=True)
combined = Workbook(); combined.remove(combined.active)
for s in sellers:
    sr = [r for r in rows if r[1] == s]
    wb = Workbook(); fill_sheet(wb.active, sr); wb.active.title = s[:31]
    wb.save(os.path.join(OUT, s.lower().replace(" ", "-") + ".xlsx"))
    fill_sheet(combined.create_sheet(s[:31]), sr)
    print("%-18s %2d lines  $%.2f" % (s, len(sr), sum(int(r[2]) * float(r[6]) for r in sr if r[6])))
combined.save(os.path.join(OUT, "all-vendors.xlsx"))
